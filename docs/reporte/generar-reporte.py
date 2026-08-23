"""Genera el reporte tecnico en Word con formato APA 7 (entregable de T7.5).

    uv run --with python-docx python docs/reporte/generar-reporte.py

Se versiona el generador y no solo el .docx para que el reporte sea
reproducible: si cambia un numero del benchmark, se edita aqui y se regenera.

Las cifras NO se inventan: salen de benchmark/results/overhead-analysis.md y de
benchmark/results/tablas-20260822-221404.md.

El documento debe quedar entre 5 y 7 paginas. Para comprobarlo de verdad, y no
estimando, se pagina con Word:

    osascript -e 'tell application "Microsoft Word"
        set d to open file name POSIX file "<ruta absoluta>" as string
        set n to (compute statistics d statistic statistic pages)
        close d saving no
        return n
    end tell'

Estado actual: 7 paginas, 1378 palabras.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

# ── Configuracion APA 7: Times New Roman 12, margenes 1", doble espacio ──────
est = DOC.styles["Normal"]
est.font.name = "Times New Roman"
est.font.size = Pt(12)
est.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = est.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
pf.space_after = Pt(0)
pf.space_before = Pt(0)

for s in DOC.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)

# ── Numero de pagina arriba a la derecha ────────────────────────────────────
hdr = DOC.sections[0].header.paragraphs[0]
hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = hdr.add_run()
for instr, attr in (("begin", None), (None, "PAGE"), ("end", None)):
    el = OxmlElement("w:fldChar") if instr else OxmlElement("w:instrText")
    if instr:
        el.set(qn("w:fldCharType"), instr)
    else:
        el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
    run._r.append(el)
hdr.runs[0].font.name = "Times New Roman"
hdr.runs[0].font.size = Pt(12)


def p(texto="", *, negrita=False, cursiva=False, align=None, sangria=True,
      espacio=False, size=12, fuente=None, sencillo=False):
    par = DOC.add_paragraph()
    par.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.SINGLE if sencillo else WD_LINE_SPACING.DOUBLE)
    if align:
        par.alignment = align
    if sangria and align is None:
        par.paragraph_format.first_line_indent = Inches(0.5)
    if espacio:
        par.paragraph_format.space_before = Pt(12)
    if texto:
        r = par.add_run(texto)
        r.bold, r.italic = negrita, cursiva
        r.font.size = Pt(size)
        if fuente:
            r.font.name = fuente
    return par


def h1(t):
    p(t, negrita=True, align=WD_ALIGN_PARAGRAPH.CENTER, espacio=True)


def h2(t):
    p(t, negrita=True, sangria=False, espacio=True)


def compacto(par):
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return par


# ════════════════════════════════════════════════════════════════════════════
# PORTADA
# ════════════════════════════════════════════════════════════════════════════
for _ in range(4):
    p()
p("Implementación de un pipeline de observabilidad con OpenTelemetry: "
  "arquitectura, correlación cross-signal y análisis de overhead",
  negrita=True, align=WD_ALIGN_PARAGRAPH.CENTER)
p()
p("Fredy Pulido, Myriam Martínez, Juan Francisco Pérez y Nicolás Torres",
  align=WD_ALIGN_PARAGRAPH.CENTER)
p("Maestría en Arquitectura de Software", align=WD_ALIGN_PARAGRAPH.CENTER)
p("MASS – OBAP20264: Observabilidad en Ambientes Productivos",
  align=WD_ALIGN_PARAGRAPH.CENTER)
p("[Nombre del docente]", align=WD_ALIGN_PARAGRAPH.CENTER)
p("25 de agosto de 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
DOC.add_page_break()

# ════════════════════════════════════════════════════════════════════════════
# CUERPO
# ════════════════════════════════════════════════════════════════════════════
p("Implementación de un pipeline de observabilidad con OpenTelemetry",
  negrita=True, align=WD_ALIGN_PARAGRAPH.CENTER)

h2("Introducción")
p("La observabilidad no verifica umbrales conocidos, sino que permite responder "
  "preguntas no anticipadas (Majors et al., 2022). El objetivo de este laboratorio no "
  "es recolectar telemetría, sino demostrar que las tres señales pueden recorrerse "
  "entre sí para investigar una petición concreta, y cuantificar el costo de esa "
  "capacidad. El sistema modela un checkout: el servicio A valida un carrito, aplica "
  "un descuento y solicita al servicio B la reserva de inventario, que consulta "
  "PostgreSQL. Se emplearon dos servicios porque el problema relevante del trazado "
  "distribuido es el salto entre procesos.")

h2("Arquitectura de la Solución")
p("La aplicación exporta telemetría por OTLP/gRPC hacia un OpenTelemetry Collector "
  "que la distribuye a tres backends especializados. Este desacople es la decisión "
  "arquitectónica central: la aplicación conoce una sola dirección y el destino se "
  "resuelve en configuración. Sin esa capa intermedia, cambiar de backend obligaría "
  "a reconstruir y redesplegar todos los servicios.")

# Tabla 1
p("Tabla 1", negrita=True, sangria=False, espacio=True)
p("Componentes del pipeline y responsabilidad de cada uno", cursiva=True, sangria=False)
t1 = DOC.add_table(rows=6, cols=3)
t1.style = "Table Grid"
filas1 = [
    ("Componente", "Señal", "Función"),
    ("OTel SDK (Python)", "Las tres", "Instrumentación automática y manual en el proceso"),
    ("OTel Collector", "Las tres", "Recepción, límite de memoria, etiquetado y reenvío"),
    ("Jaeger", "Trazas", "Almacenamiento y visualización en cascada"),
    ("Prometheus", "Métricas", "Series temporales y exemplars"),
    ("Loki", "Logs", "Indexación por etiquetas, no por texto"),
]
for i, fila in enumerate(filas1):
    for j, celda in enumerate(fila):
        c = t1.cell(i, j)
        c.text = celda
        par = c.paragraphs[0]
        compacto(par)
        par.runs[0].font.size = Pt(10)
        par.runs[0].font.name = "Times New Roman"
        if i == 0:
            par.runs[0].bold = True
p("Nota. El Collector aplica los procesadores en el orden memory_limiter → resource "
  "→ resourcedetection → filter → batch.", sangria=False, size=10, espacio=True, sencillo=True)

h2("Decisiones de Diseño")
p("Instrumentación automática y manual. La automática se activa con el comando "
  "opentelemetry-instrument, sin envolver el código: cubre cada petición HTTP, cada "
  "consulta SQL y la propagación de contexto W3C (W3C, 2021). Su límite es que "
  "desconoce el dominio, pues observa una sentencia UPDATE y no una reserva de "
  "inventario; por ello se añadieron tres spans de negocio con atributos propios.")
p("Orden de los procesadores y cardinalidad. El limitador de memoria va primero "
  "porque su función es rechazar datos bajo presión: situado tras el agrupamiento, el "
  "Collector ya habría consumido la memoria que intenta proteger, y uno que termina "
  "por falta de memoria deja al sistema ciego durante un incidente. El agrupamiento va "
  "último, justo antes de la salida de red. Por su parte, los identificadores de "
  "carrito se registran como atributos de span y nunca como etiquetas de métrica, pues "
  "una etiqueta de cardinalidad alta genera una serie temporal por valor y agota la "
  "memoria del sistema de métricas.")

h2("Correlación Cross-Signal")
p("La correlación se verificó sobre una misma petición, identificada por el trazo "
  "d0d3061140d349621409832fbf158bce. Prometheus almacenó un exemplar de 809 ms que "
  "referencia ese identificador; Jaeger contiene la traza con 15 spans y 810,29 ms "
  "repartidos entre ambos servicios; y Loki devuelve siete líneas de registro de los "
  "dos servicios al filtrar por ese valor. La coincidencia entre el exemplar y la "
  "duración real confirma que la cadena opera de extremo a extremo. Ello exige cinco "
  "eslabones correctamente configurados, desde el filtro de exemplars en el SDK hasta "
  "los enlaces entre orígenes de datos en Grafana, y cada uno falla en silencio.")

# Figura 1
p("Figura 1", negrita=True, sangria=False, espacio=True)
p("Traza distribuida con spans automáticos y de negocio en un mismo árbol",
  cursiva=True, sangria=False)
DOC.add_picture("docs/evidencias/R3-01-traza.png", width=Inches(3.9))
DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Nota. Los spans checkout.validate_cart e inventory.reserve_stock son manuales; "
  "SELECT y UPDATE los genera la instrumentación automática de psycopg2 y aparecen "
  "anidados bajo el span de negocio.", sangria=False, size=10, sencillo=True)

h2("Análisis de Overhead")
p("Se compararon dos escenarios con k6 —SDK desactivado frente a instrumentación "
  "completa— bajo 50 usuarios virtuales durante 420 segundos, con tres corridas por "
  "escenario y descarte de la primera como calentamiento. Las seis resultaron "
  "válidas, con cero errores inesperados.")

# Tabla 2
p("Tabla 2", negrita=True, sangria=False, espacio=True)
p("Latencia y rendimiento por escenario", cursiva=True, sangria=False)
t2 = DOC.add_table(rows=5, cols=4)
t2.style = "Table Grid"
filas2 = [
    ("Métrica", "Sin instrumentar", "Instrumentado", "Diferencia"),
    ("Latencia p50", "135,5 ms", "167,5 ms", "+32,0 ms (23,6 %)"),
    ("Latencia p95", "199,5 ms", "257,0 ms", "+57,5 ms (28,8 %)"),
    ("Latencia p99", "245,5 ms", "329,0 ms", "+83,5 ms (34,0 %)"),
    ("Rendimiento", "336,6 sol/s", "270,1 sol/s", "−19,7 %"),
]
for i, fila in enumerate(filas2):
    for j, celda in enumerate(fila):
        c = t2.cell(i, j)
        c.text = celda
        par = c.paragraphs[0]
        compacto(par)
        par.runs[0].font.size = Pt(10)
        par.runs[0].font.name = "Times New Roman"
        if i == 0:
            par.runs[0].bold = True
p("Nota. Media de dos corridas. La desviación entre corridas fue de 5,3 % y 1,7 % "
  "sobre el p95, frente a una diferencia entre escenarios de 28,8 %.",
  sangria=False, size=10, espacio=True, sencillo=True)

p("Los cuatro valores de la Tabla 2 no son hallazgos independientes. La prueba opera "
  "en lazo cerrado con concurrencia fija, régimen donde rige la ley de Little (Little, "
  "1961): el tiempo de respuesta equivale al número de usuarios dividido entre el "
  "rendimiento. Esa relación predice un aumento de latencia del 24,6 % y se midió "
  "23,6 %, lo que confirma que el alza de latencia y la caída de rendimiento son el "
  "mismo fenómeno visto desde dos ángulos.")
p("El consumo de procesador exigió otro tratamiento. En términos absolutos resultaba "
  "engañoso, pues el servicio A registró menos uso al estar instrumentado; no es un "
  "ahorro, sino menos peticiones atendidas, ya que ambos escenarios operan saturados. "
  "Normalizar el consumo por el rendimiento entrega la magnitud comparable.")

# Tabla 3
p("Tabla 3", negrita=True, sangria=False, espacio=True)
p("Consumo de procesador por petición", cursiva=True, sangria=False)
t3 = DOC.add_table(rows=4, cols=4)
t3.style = "Table Grid"
filas3 = [
    ("Capa", "Sin instrumentar", "Instrumentado", "Diferencia"),
    ("Aplicación y base de datos", "0,7937", "0,9452", "+19,1 %"),
    ("Backends de telemetría", "0,0042", "0,0797", "—"),
    ("Total del sistema", "0,7978", "1,0249", "+28,5 %"),
]
for i, fila in enumerate(filas3):
    for j, celda in enumerate(fila):
        c = t3.cell(i, j)
        c.text = celda
        par = c.paragraphs[0]
        compacto(par)
        par.runs[0].font.size = Pt(10)
        par.runs[0].font.name = "Times New Roman"
        if i == 0:
            par.runs[0].bold = True
p("Nota. Unidad: porcentaje de procesador dividido entre solicitudes por segundo. Del "
  "sobrecosto total, 67 % corresponde a la aplicación y 33 % a los backends. El "
  "consumo adicional de memoria fue de 4,6 MB y 4,3 MB por servicio.",
  sangria=False, size=10, espacio=True, sencillo=True)

p("El costo transferible es el 28,5 % de procesador por petición, no el 34 % del "
  "percentil 99: este último depende de que el servicio B operara al 98,8 % de "
  "procesador. En un sistema con holgura, el mismo incremento de tiempo de servicio "
  "produce un alza de latencia mucho menor, al no existir una cola que la amplifique.")

p("La primera ejecución resultó inválida por miles de errores. Prometheus descartó "
  "la hipótesis de agotamiento de inventario al no registrar error de cliente alguno, "
  "y los registros revelaron la causa real: el conjunto de conexiones admitía cinco "
  "frente a los cuarenta hilos con que el servidor atiende funciones síncronas, y "
  "empleaba una implementación no segura para hilos. Corregido el defecto, la medición "
  "se repitió. Una prueba de humo secuencial jamás lo habría detectado, lo que respalda "
  "medir bajo carga y no solo comprobar que los extremos responden.")

# Figura 2
p("Figura 2", negrita=True, sangria=False, espacio=True)
p("Exemplar que enlaza un punto de la métrica de latencia con su traza",
  cursiva=True, sangria=False)
DOC.add_picture("docs/evidencias/R3-03-exemplar.png", width=Inches(3.6))
DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Nota. El cuadro emergente muestra el identificador de traza asociado al valor de "
  "809 ms y el enlace directo hacia Jaeger.", sangria=False, size=10, sencillo=True)

h2("Conclusiones y Recomendaciones")
p("La instrumentación completa costó 28,5 % de procesador por petición y 4,5 MB de "
  "memoria por servicio, a cambio de poder investigar cualquier petición individual a "
  "través de las tres señales. El costo de memoria es despreciable y acotado; el de "
  "procesador justifica muestreo en alto volumen.")
p("La distribución del sobrecosto orienta esa estrategia en sentido contrario al "
  "habitual: como dos tercios se consumen dentro de la aplicación, el muestreo en el "
  "Collector solo reduciría el tercio restante, pues el gasto ya ocurrió antes de que "
  "el Collector recibiera los datos. La palanca efectiva es el muestreo en origen, que "
  "evita crear el span; su contrapartida es la pérdida de trazas de error, ya que la "
  "decisión se toma al inicio, y se mitiga conservando métricas y registros íntegros, "
  "señales mucho más económicas (Beyer et al., 2018). Se recomienda muestreo íntegro "
  "en servicios críticos o de bajo volumen, y entre 10 % y 20 % en origen para "
  "servicios de alto volumen sin holgura de procesador.")

# ════════════════════════════════════════════════════════════════════════════
# REFERENCIAS
# ════════════════════════════════════════════════════════════════════════════
DOC.add_page_break()
p("Referencias", negrita=True, align=WD_ALIGN_PARAGRAPH.CENTER)
refs = [
    "Beyer, B., Jones, C., Petoff, J., & Murphy, N. R. (2016). Site reliability "
    "engineering: How Google runs production systems. O'Reilly Media.",
    "Beyer, B., Murphy, N. R., Rensin, D. K., Kawahara, K., & Thorne, S. (2018). "
    "The site reliability workbook: Practical ways to implement SRE. O'Reilly Media.",
    "Little, J. D. C. (1961). A proof for the queuing formula: L = λW. "
    "Operations Research, 9(3), 383–387. https://doi.org/10.1287/opre.9.3.383",
    "Majors, C., Fong-Jones, L., & Miranda, G. (2022). Observability engineering: "
    "Achieving production excellence. O'Reilly Media.",
    "OpenTelemetry Authors. (2025). OpenTelemetry documentation. Cloud Native "
    "Computing Foundation. https://opentelemetry.io/docs/",
    "Sigelman, B. H., Barroso, L. A., Burrows, M., Stephenson, P., Plakal, M., "
    "Beaver, D., Jaspan, S., & Shanbhag, C. (2010). Dapper, a large-scale distributed "
    "systems tracing infrastructure (Technical Report dapper-2010-1). Google.",
    "World Wide Web Consortium. (2021). Trace context (W3C Recommendation). "
    "https://www.w3.org/TR/trace-context/",
]
for r in refs:
    par = p(r, sangria=False)
    par.paragraph_format.left_indent = Inches(0.5)
    par.paragraph_format.first_line_indent = Inches(-0.5)

DOC.save("docs/reporte/Reporte-Tecnico-OTel-MASS-OBAP20264.docx")
print("documento generado")
